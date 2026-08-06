"""
Sales Manager — Database migration and API endpoint code.

== DATABASE MIGRATION ==
Run this SQL or add to database.py _ensure_tables():
"""

DB_MIGRATION_SQL = """
-- Vendor sales
IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'vendor_sales')
CREATE TABLE vendor_sales (
    id INT IDENTITY(1,1) PRIMARY KEY,
    vendor NVARCHAR(200) NOT NULL,
    name NVARCHAR(300) NOT NULL,
    status NVARCHAR(20) NOT NULL DEFAULT 'pending',
    currency NVARCHAR(10) NOT NULL DEFAULT 'USD',
    fx_rate_at_creation DECIMAL(10,6) NOT NULL DEFAULT 1.0,
    start_at DATETIME2 NULL,
    end_at DATETIME2 NULL,
    notes NVARCHAR(MAX) NULL,
    created_at DATETIME2 NOT NULL DEFAULT GETUTCDATE(),
    activated_at DATETIME2 NULL,
    reverted_at DATETIME2 NULL
);

-- Vendor sale items
IF NOT EXISTS (SELECT * FROM sys.tables WHERE name = 'vendor_sale_items')
CREATE TABLE vendor_sale_items (
    id INT IDENTITY(1,1) PRIMARY KEY,
    sale_id INT NOT NULL,
    sku NVARCHAR(100) NOT NULL,
    shopify_variant_id NVARCHAR(100) NULL,
    shopify_product_id NVARCHAR(100) NULL,
    original_price DECIMAL(12,2) NOT NULL DEFAULT 0,
    original_compare_at DECIMAL(12,2) NULL,
    sale_price_foreign DECIMAL(12,2) NOT NULL DEFAULT 0,
    sale_price_cad DECIMAL(12,2) NOT NULL DEFAULT 0,
    discount_pct DECIMAL(8,2) NULL,
    cost_cad DECIMAL(12,2) NULL,
    sale_margin_pct DECIMAL(8,2) NULL,
    status NVARCHAR(20) NOT NULL DEFAULT 'pending',
    CONSTRAINT FK_sale_items_sale FOREIGN KEY (sale_id) REFERENCES vendor_sales(id)
);

IF NOT EXISTS (SELECT * FROM sys.indexes WHERE name = 'IX_vendor_sale_items_sale_id')
    CREATE INDEX IX_vendor_sale_items_sale_id ON vendor_sale_items(sale_id);

IF NOT EXISTS (SELECT * FROM sys.indexes WHERE name = 'IX_vendor_sale_items_sku')
    CREATE INDEX IX_vendor_sale_items_sku ON vendor_sale_items(sku);

IF NOT EXISTS (SELECT * FROM sys.indexes WHERE name = 'IX_vendor_sales_status')
    CREATE INDEX IX_vendor_sales_status ON vendor_sales(status);
"""


# ─────────────────────────────────────────────────────────────────
# SHOPIFY CLIENT — add this method to shopify_client.py
# ─────────────────────────────────────────────────────────────────

SHOPIFY_METHOD = '''
    async def update_variant_price_with_compare(self, product_id: str,
                                                  variant_id: str,
                                                  price: float,
                                                  compare_at_price: float = None) -> Dict:
        """Update a variant's price and compareAtPrice in Shopify."""
        mutation = """
        mutation($productId: ID!, $variants: [ProductVariantsBulkInput!]!) {
            productVariantsBulkUpdate(productId: $productId, variants: $variants) {
                productVariants {
                    id
                    price
                    compareAtPrice
                }
                userErrors {
                    field
                    message
                }
            }
        }
        """
        variant_input = {
            "id": variant_id,
            "price": str(round(price, 2)),
        }
        if compare_at_price is not None:
            variant_input["compareAtPrice"] = str(round(compare_at_price, 2))
        else:
            variant_input["compareAtPrice"] = None

        variables = {
            "productId": product_id,
            "variants": [variant_input],
        }
        data = await self._query(mutation, variables)
        result = data["productVariantsBulkUpdate"]
        errors = result.get("userErrors", [])
        if errors:
            raise Exception(f"Price update error: {errors[0].get('message', 'Unknown')}")
        return result
'''


# ─────────────────────────────────────────────────────────────────
# API ENDPOINTS — add to main.py
# ─────────────────────────────────────────────────────────────────

API_ENDPOINTS = '''
# ─── SALES MANAGER ──────────────────────────────────────────────

from .sales_manager import (
    create_sale, list_sales, get_sale, update_sale, delete_sale,
    preview_sale_pricelist, confirm_sale_items,
    activate_sale, revert_sale, exclude_item_from_sale,
    check_scheduled_sales,
)


class CreateSaleRequest(BaseModel):
    vendor: str
    name: str
    currency: str = "USD"
    fx_rate: float = 1.0
    start_at: Optional[str] = None
    end_at: Optional[str] = None
    notes: Optional[str] = None


class UpdateSaleRequest(BaseModel):
    name: Optional[str] = None
    status: Optional[str] = None
    start_at: Optional[str] = None
    end_at: Optional[str] = None
    notes: Optional[str] = None


class SalePreviewRequest(BaseModel):
    sku_col: str
    price_col: str
    price_currency: str = "USD"
    fx_rate: float = 1.0


class SaleConfirmRequest(BaseModel):
    items: List[dict]


class SaleExcludeRequest(BaseModel):
    sku: str


@app.get("/api/sales")
async def api_list_sales(
    vendor: Optional[str] = None,
    status: Optional[str] = None,
    token: str = Depends(verify_token),
):
    """List all sales, optionally filtered by vendor or status."""
    try:
        return list_sales(db, vendor, status)
    except Exception as e:
        logger.error(f"List sales error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sales")
async def api_create_sale(data: CreateSaleRequest, token: str = Depends(verify_token)):
    """Create a new vendor sale."""
    try:
        return create_sale(db, data.vendor, data.name, data.currency,
                          data.fx_rate, data.start_at, data.end_at, data.notes)
    except Exception as e:
        logger.error(f"Create sale error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/sales/{sale_id}")
async def api_get_sale(sale_id: int, token: str = Depends(verify_token)):
    """Get a sale with all its items."""
    try:
        result = get_sale(db, sale_id)
        if not result:
            raise HTTPException(status_code=404, detail="Sale not found")
        return result
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Get sale error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.put("/api/sales/{sale_id}")
async def api_update_sale(sale_id: int, data: UpdateSaleRequest,
                          token: str = Depends(verify_token)):
    """Update sale metadata."""
    try:
        updates = {k: v for k, v in data.dict().items() if v is not None}
        return update_sale(db, sale_id, updates)
    except Exception as e:
        logger.error(f"Update sale error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.delete("/api/sales/{sale_id}")
async def api_delete_sale(sale_id: int, token: str = Depends(verify_token)):
    """Delete a pending or cancelled sale."""
    try:
        return delete_sale(db, sale_id)
    except Exception as e:
        logger.error(f"Delete sale error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sales/{sale_id}/upload")
async def api_upload_sale_pricelist(
    sale_id: int,
    file: UploadFile = File(...),
    sku_col: str = Query("SKU"),
    price_col: str = Query("Sale Price"),
    price_currency: str = Query("USD"),
    fx_rate: float = Query(1.0),
    token: str = Depends(verify_token),
):
    """Upload a sale pricelist and get a preview of changes."""
    try:
        raw = await file.read()
        filename = file.filename or ""

        if filename.lower().endswith(('.xlsx', '.xls')):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise HTTPException(400, "Empty spreadsheet")
            headers = [str(h or '').strip() for h in rows[0]]
            csv_lines = [','.join(headers)]
            for row in rows[1:]:
                csv_lines.append(','.join(str(c or '') for c in row))
            csv_text = '\\n'.join(csv_lines)
        else:
            try:
                csv_text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                csv_text = raw.decode('latin-1')

        return preview_sale_pricelist(db, sale_id, csv_text,
                                      sku_col, price_col,
                                      price_currency, fx_rate)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Sale pricelist upload error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sales/{sale_id}/confirm")
async def api_confirm_sale(sale_id: int, data: SaleConfirmRequest,
                           token: str = Depends(verify_token)):
    """Save previewed items to the sale (does not activate)."""
    try:
        return confirm_sale_items(db, sale_id, data.items)
    except Exception as e:
        logger.error(f"Confirm sale error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sales/{sale_id}/activate")
async def api_activate_sale(sale_id: int, token: str = Depends(verify_token)):
    """Activate a sale: apply sale prices to Shopify."""
    try:
        return await activate_sale(db, shopify_client, sale_id)
    except Exception as e:
        logger.error(f"Activate sale error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sales/{sale_id}/revert")
async def api_revert_sale(sale_id: int, token: str = Depends(verify_token)):
    """Revert a sale: restore original prices."""
    try:
        return await revert_sale(db, shopify_client, sale_id)
    except Exception as e:
        logger.error(f"Revert sale error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sales/{sale_id}/exclude")
async def api_exclude_sale_item(sale_id: int, data: SaleExcludeRequest,
                                 token: str = Depends(verify_token)):
    """Remove a single SKU from an active sale."""
    try:
        return await exclude_item_from_sale(db, shopify_client, sale_id, data.sku)
    except Exception as e:
        logger.error(f"Exclude sale item error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sales/check-schedule")
async def api_check_sale_schedule(token: str = Depends(verify_token)):
    """Check for sales that need to be activated or reverted based on schedule."""
    try:
        actions = check_scheduled_sales(db)
        results = []
        for action in actions:
            if action["action"] == "activate":
                r = await activate_sale(db, shopify_client, action["sale_id"])
                results.append({**action, "result": r})
            elif action["action"] == "revert":
                r = await revert_sale(db, shopify_client, action["sale_id"])
                results.append({**action, "result": r})
        return {"actions_processed": len(results), "results": results}
    except Exception as e:
        logger.error(f"Schedule check error: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))
'''


# ─────────────────────────────────────────────────────────────────
# API.JS — add to frontend/src/api.js
# ─────────────────────────────────────────────────────────────────

API_JS = '''
// ── Sales Manager ─────────────────────────────────────────────

export async function listSales(vendor, status) {
  const params = new URLSearchParams();
  if (vendor) params.append('vendor', vendor);
  if (status) params.append('status', status);
  const qs = params.toString();
  return apiFetch('/sales' + (qs ? '?' + qs : ''));
}

export async function createSale(data) {
  return apiFetch('/sales', { method: 'POST', body: JSON.stringify(data) });
}

export async function getSale(saleId) {
  return apiFetch('/sales/' + saleId);
}

export async function updateSale(saleId, data) {
  return apiFetch('/sales/' + saleId, { method: 'PUT', body: JSON.stringify(data) });
}

export async function deleteSale(saleId) {
  return apiFetch('/sales/' + saleId, { method: 'DELETE' });
}

export async function uploadSalePricelist(saleId, file, skuCol, priceCol, currency, fxRate) {
  const token = localStorage.getItem('tc_planner_token') || '';
  const formData = new FormData();
  formData.append('file', file);
  const params = new URLSearchParams({
    sku_col: skuCol, price_col: priceCol,
    price_currency: currency, fx_rate: String(fxRate),
  });
  const response = await fetch(API_BASE + '/sales/' + saleId + '/upload?' + params.toString(), {
    method: 'POST',
    headers: { Authorization: 'Bearer ' + token },
    body: formData,
  });
  if (!response.ok) {
    const error = await response.json().catch(() => ({ detail: 'Upload failed' }));
    throw new Error(error.detail || 'HTTP ' + response.status);
  }
  return response.json();
}

export async function confirmSale(saleId, items) {
  return apiFetch('/sales/' + saleId + '/confirm', {
    method: 'POST', body: JSON.stringify({ items }),
  });
}

export async function activateSale(saleId) {
  return apiFetch('/sales/' + saleId + '/activate', { method: 'POST' });
}

export async function revertSale(saleId) {
  return apiFetch('/sales/' + saleId + '/revert', { method: 'POST' });
}

export async function excludeSaleItem(saleId, sku) {
  return apiFetch('/sales/' + saleId + '/exclude', {
    method: 'POST', body: JSON.stringify({ sku }),
  });
}

export async function checkSaleSchedule() {
  return apiFetch('/sales/check-schedule', { method: 'POST' });
}
'''
